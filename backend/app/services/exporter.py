"""
PDF and Excel export of CVE alerts for a given customer.
"""
import io
from datetime import datetime
from sqlalchemy.orm import Session, joinedload
from app.models.cve import CVEAlert
from app.models.customer import Customer


def _get_cves_for_customer(customer: Customer, db: Session):
    alerts = (
        db.query(CVEAlert)
        .filter(CVEAlert.customer_id == customer.id)
        .options(joinedload(CVEAlert.cve))
        .order_by(CVEAlert.created_at.desc())
        .all()
    )
    return [a.cve for a in alerts]


def generate_pdf(customer: Customer, db: Session) -> bytes:
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib import colors
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
    from reportlab.lib.units import mm

    cves = _get_cves_for_customer(customer, db)
    buffer = io.BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=landscape(A4), leftMargin=15*mm, rightMargin=15*mm,
                            topMargin=15*mm, bottomMargin=15*mm)
    styles = getSampleStyleSheet()
    elements = []

    title_style = ParagraphStyle("title", parent=styles["Heading1"], fontSize=14)
    elements.append(Paragraph(f"CVE Report — {customer.name}", title_style))
    elements.append(Paragraph(f"Generated: {datetime.utcnow().strftime('%Y-%m-%d %H:%M UTC')}", styles["Normal"]))
    elements.append(Spacer(1, 8*mm))

    severity_colors = {
        "CRITICAL": colors.HexColor("#c0392b"),
        "HIGH": colors.HexColor("#e67e22"),
        "MEDIUM": colors.HexColor("#f39c12"),
        "LOW": colors.HexColor("#27ae60"),
    }

    data = [["CVE ID", "Severity", "CVSS", "Published", "Description"]]
    for cve in cves:
        desc = (cve.description or "")[:120] + ("…" if cve.description and len(cve.description) > 120 else "")
        data.append([
            cve.cve_id,
            cve.severity or "—",
            str(cve.cvss_score or "—"),
            cve.published_at.strftime("%Y-%m-%d") if cve.published_at else "—",
            desc,
        ])

    col_widths = [60*mm, 25*mm, 20*mm, 25*mm, 130*mm]
    table = Table(data, colWidths=col_widths, repeatRows=1)
    table_style = [
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1a3a5c")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("FONTSIZE", (0, 0), (-1, -1), 8),
        ("GRID", (0, 0), (-1, -1), 0.3, colors.grey),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#f8f9fa")]),
    ]
    for i, cve in enumerate(cves, start=1):
        if cve.severity in severity_colors:
            table_style.append(("TEXTCOLOR", (1, i), (1, i), severity_colors[cve.severity]))
            table_style.append(("FONTNAME", (1, i), (1, i), "Helvetica-Bold"))
    table.setStyle(TableStyle(table_style))
    elements.append(table)

    doc.build(elements)
    return buffer.getvalue()


def generate_excel(customer: Customer, db: Session) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment

    cves = _get_cves_for_customer(customer, db)
    wb = Workbook()
    ws = wb.active
    ws.title = "CVEs"

    headers = ["CVE ID", "Severity", "CVSS Score", "CVSS Version", "Published", "Source", "URL", "Description"]
    header_fill = PatternFill(start_color="1a3a5c", end_color="1a3a5c", fill_type="solid")
    header_font = Font(color="FFFFFF", bold=True)

    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center")

    severity_fills = {
        "CRITICAL": PatternFill(start_color="FFCCCC", end_color="FFCCCC", fill_type="solid"),
        "HIGH":     PatternFill(start_color="FFE0CC", end_color="FFE0CC", fill_type="solid"),
        "MEDIUM":   PatternFill(start_color="FFF3CC", end_color="FFF3CC", fill_type="solid"),
        "LOW":      PatternFill(start_color="CCFFCC", end_color="CCFFCC", fill_type="solid"),
    }

    for row, cve in enumerate(cves, 2):
        ws.cell(row=row, column=1, value=cve.cve_id)
        ws.cell(row=row, column=2, value=cve.severity or "")
        ws.cell(row=row, column=3, value=cve.cvss_score)
        ws.cell(row=row, column=4, value=cve.cvss_version or "")
        ws.cell(row=row, column=5, value=cve.published_at.strftime("%Y-%m-%d") if cve.published_at else "")
        ws.cell(row=row, column=6, value=cve.source or "")
        ws.cell(row=row, column=7, value=cve.source_url or "")
        ws.cell(row=row, column=8, value=cve.description or "")

        if cve.severity in severity_fills:
            for col in range(1, 9):
                ws.cell(row=row, column=col).fill = severity_fills[cve.severity]

    ws.column_dimensions["A"].width = 18
    ws.column_dimensions["B"].width = 12
    ws.column_dimensions["H"].width = 80

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()
